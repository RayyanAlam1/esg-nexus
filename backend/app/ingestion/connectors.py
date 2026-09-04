"""Ingestion connectors (adapter pattern). Each connector turns an input (bytes/file) into row dicts.

Register a new connector with @connector("kind") — see docs/guides/add-connector.md.
"""

from __future__ import annotations

import csv
import io
import json
import re
from typing import Callable

CONNECTORS: dict[str, Callable[[bytes, str], list[dict]]] = {}


def connector(kind: str):
    def deco(fn):
        CONNECTORS[kind] = fn
        return fn

    return deco


def _norm(v):
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"-?[\d,]+(\.\d+)?", s):
            try:
                return float(s.replace(",", ""))
            except ValueError:
                return s
        return s
    return v


@connector("csv")
def read_csv(data: bytes, file_name: str = "") -> list[dict]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    return [{(k or "").strip(): _norm(v) for k, v in row.items()} for row in reader]


@connector("json")
def read_json(data: bytes, file_name: str = "") -> list[dict]:
    obj = json.loads(data.decode("utf-8"))
    rows = obj if isinstance(obj, list) else obj.get("rows") or obj.get("data") or []
    return [{k: _norm(v) for k, v in r.items()} for r in rows if isinstance(r, dict)]


@connector("excel")
def read_excel(data: bytes, file_name: str = "") -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    return [{header[i]: _norm(v) for i, v in enumerate(r) if i < len(header)} for r in rows[1:] if any(v is not None for v in r)]


@connector("pdf")
def read_pdf(data: bytes, file_name: str = "") -> list[dict]:
    """Document connector: extracts page text plus candidate 'label: number' facts for the ESG Data Agent to map."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    rows = []
    fact_re = re.compile(r"([A-Za-z][A-Za-z &/\-\(\)]{4,60}?)\s+(-?\d[\d,]*(?:\.\d+)?)\s*(%|GJ|tCO2e|ML|t|kt|KT|hours|MW|kW|PKR|USD|mn)?", re.I)
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        rows.append({"page": i, "kind": "page_text", "text": text[:8000]})
        for m in fact_re.finditer(text):
            label, num, unit = m.group(1).strip(), m.group(2), (m.group(3) or "").strip()
            try:
                rows.append({"page": i, "kind": "candidate_fact", "label": label, "value": float(num.replace(",", "")), "unit": unit})
            except ValueError:
                continue
    return rows


@connector("text")
def read_text(data: bytes, file_name: str = "") -> list[dict]:
    return [{"kind": "text", "text": data.decode("utf-8", errors="replace")[:20000]}]


def detect_kind(file_name: str) -> str:
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    return {"csv": "csv", "json": "json", "xlsx": "excel", "xlsm": "excel", "xls": "excel", "pdf": "pdf", "txt": "text", "md": "text"}.get(ext, "csv")
