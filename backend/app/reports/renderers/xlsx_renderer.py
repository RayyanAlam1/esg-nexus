"""Excel data appendix (openpyxl): metrics sheet, per-section tables, evidence references."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="0F2A44")


def _sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title[:31])
    ws.append(columns)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append([("" if v is None else v) for v in r])
    for i, _ in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = 22 if i > 1 else 40
    ws.freeze_panes = "A2"
    return ws


def render(payload: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(
        wb,
        "Metrics",
        ["Code", "Metric", "Pillar", "Unit", payload["period_code"], "Prior", "Status", "Kind", "Assurance", "Evidence"],
        [
            [r["code"], r["name"], r["pillar"], r["unit"], r["value"], r["previous"], r["status"], r["kind"], r["assurance"], ", ".join(r["evidence"])]
            for r in payload["metric_rows"]
        ],
    )
    for s in payload["sections"]:
        for i, t in enumerate(s.get("tables") or []):
            _sheet(wb, f"{s['code'][:20]}_{i + 1}", t["columns"], t["rows"])
    _sheet(
        wb,
        "Info",
        ["Field", "Value"],
        [
            ["Title", payload["title"]],
            ["Organization", payload["organization"]],
            ["Period", payload["period"]],
            ["Frameworks", ", ".join(payload["frameworks"])],
            ["Generated", payload["generated_at"]],
            ["Status", payload["status"]],
            ["Watermark", payload.get("watermark") or ""],
        ],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
